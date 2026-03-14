from __future__ import annotations

import csv
import json
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def _infer_type(values: list[str]) -> str:
    material = [value for value in values if value != ""]
    if not material:
        return "empty"
    numeric = 0
    for value in material:
        try:
            float(value)
            numeric += 1
        except ValueError:
            continue
    if numeric == len(material):
        return "number"
    return "string"


def _select_header_index(rows: list[list[str]]) -> int:
    best_index = 0
    best_score = -1.0
    for index, row in enumerate(rows[:8]):
        non_empty = [cell for cell in row if cell]
        if not non_empty:
            continue
        text_like = sum(1 for cell in non_empty if _infer_type([cell]) == "string")
        unique_count = len(set(non_empty))
        score = (len(non_empty) * 2) + text_like + (unique_count * 0.5)
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


@dataclass(frozen=True)
class ColumnProfile:
    column_name: str
    normalized_name: str
    inferred_type: str
    unique_ratio: float
    null_ratio: float
    sample_values: list[str]
    top_values: list[str]


@dataclass(frozen=True)
class TableProfile:
    source_path: Path
    kind: str
    logical_name: str
    headers: list[str]
    columns: list[ColumnProfile]
    sample_rows: list[dict[str, str]]
    data_row_count: int
    sheet_name: str | None = None
    sheet_index: int | None = None
    is_primary_candidate: bool = False

    def to_record(self) -> dict[str, Any]:
        return {
            "source_path": str(self.source_path),
            "scope_key": self.sheet_name or "",
            "kind": self.kind,
            "logical_name": self.logical_name,
            "sheet_name": self.sheet_name,
            "sheet_index": self.sheet_index,
            "is_primary_candidate": 1 if self.is_primary_candidate else 0,
            "headers_json": json.dumps(self.headers, ensure_ascii=False),
            "sample_rows_json": json.dumps(self.sample_rows, ensure_ascii=False),
            "data_row_count": self.data_row_count,
        }


def table_profile_from_payload(payload: dict[str, Any]) -> TableProfile:
    columns = [
        ColumnProfile(
            column_name=str(item["column_name"]),
            normalized_name=str(item.get("normalized_name", item["column_name"])).casefold(),
            inferred_type=str(item.get("inferred_type", "string")),
            unique_ratio=float(item.get("unique_ratio", 0.0)),
            null_ratio=float(item.get("null_ratio", 0.0)),
            sample_values=[str(value) for value in item.get("sample_values", [])],
            top_values=[str(value) for value in item.get("top_values", [])],
        )
        for item in payload.get("columns", [])
    ]
    return TableProfile(
        source_path=Path(str(payload["source_path"])),
        kind=str(payload["kind"]),
        logical_name=str(payload["logical_name"]),
        headers=[str(value) for value in payload.get("headers", payload.get("headers_json") and json.loads(payload["headers_json"]) or [])],
        columns=columns,
        sample_rows=[{str(key): _normalize_text(value) for key, value in row.items()} for row in payload.get("sample_rows", payload.get("sample_rows_json") and json.loads(payload["sample_rows_json"]) or [])],
        data_row_count=int(payload.get("data_row_count", 0)),
        sheet_name=str(payload["sheet_name"]) if payload.get("sheet_name") is not None else None,
        sheet_index=int(payload["sheet_index"]) if payload.get("sheet_index") is not None else None,
        is_primary_candidate=bool(payload.get("is_primary_candidate", False)),
    )


def _build_column_profiles(headers: list[str], rows: list[dict[str, str]], alias_headers: list[str] | None = None) -> list[ColumnProfile]:
    columns: list[ColumnProfile] = []
    row_count = max(len(rows), 1)
    for header in headers:
        alias = ""
        if alias_headers:
            try:
                alias = alias_headers[headers.index(header)]
            except ValueError:
                alias = ""
        values = [_normalize_text(row.get(header, "")) for row in rows]
        material = [value for value in values if value != ""]
        counter = Counter(material)
        columns.append(
            ColumnProfile(
                column_name=header,
                normalized_name=f"{header} {alias}".strip().casefold(),
                inferred_type=_infer_type(values),
                unique_ratio=0.0 if not material else round(len(set(material)) / len(material), 4),
                null_ratio=round((row_count - len(material)) / row_count, 4),
                sample_values=material[:5],
                top_values=[value for value, _ in counter.most_common(5)],
            )
        )
    return columns


def profile_csv_file(path: Path) -> TableProfile:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))

    display_row = rows[0] if rows else []
    header_row = rows[2] if len(rows) >= 3 else rows[0]
    headers = [_normalize_text(value) for value in header_row]
    alias_headers = [_normalize_text(value) for value in display_row]
    data_rows = rows[3:] if len(rows) > 3 else rows[1:]
    sample_rows: list[dict[str, str]] = []
    for row in data_rows[:10]:
        mapped = {headers[index]: _normalize_text(row[index]) if index < len(row) else "" for index in range(len(headers))}
        sample_rows.append(mapped)

    return TableProfile(
        source_path=path,
        kind="csv",
        logical_name=path.stem,
        headers=headers,
        columns=_build_column_profiles(headers, sample_rows, alias_headers=alias_headers),
        sample_rows=sample_rows,
        data_row_count=len(data_rows),
    )


def profile_workbook(path: Path) -> list[TableProfile]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    profiles: list[TableProfile] = []
    for sheet_index, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True, max_row=12)]
        normalized_rows = [[_normalize_text(value) for value in row] for row in rows if any(_normalize_text(value) for value in row)]
        header_index = _select_header_index(normalized_rows) if normalized_rows else 0
        headers = normalized_rows[header_index] if normalized_rows else []
        data_rows = normalized_rows[header_index + 1 :]
        sample_rows: list[dict[str, str]] = []
        for row in data_rows[:10]:
            mapped = {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))}
            sample_rows.append(mapped)

        profiles.append(
            TableProfile(
                source_path=path,
                kind="xlsx",
                logical_name=path.stem,
                sheet_name=sheet_name,
                sheet_index=sheet_index,
                is_primary_candidate=sheet_index == 0,
                headers=headers,
                columns=_build_column_profiles(headers, sample_rows, alias_headers=headers),
                sample_rows=sample_rows,
                data_row_count=max(sheet.max_row - 1, 0),
            )
        )

    return profiles
